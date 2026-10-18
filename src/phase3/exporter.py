"""
MANDOROS Phase 3: Research Dossier & Multi-Sheet Excel Exporter.
Generates comprehensive 7-sheet research packages (.xlsx) and publication-grade Markdown dossiers.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Dict, Any, List, Optional

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.phase3.models import Phase3Dossier


class Phase3DossierExporter:
    """
    Exports the complete 5-pillar Phase 3 research synthesis into:
    1. A publication-grade 7-sheet audit Excel workbook (.xlsx).
    2. A standalone structured Markdown dossier (.md).
    """

    @classmethod
    def export_excel(
        cls,
        dossier: Phase3Dossier,
        output_dir: str = "outputs"
    ) -> str:
        """Generates the professional 7-sheet Phase 3 Excel package."""
        os.makedirs(output_dir, exist_ok=True)
        safe_topic = dossier.target_topic.replace(" ", "_").replace("/", "_")[:30]
        excel_path = os.path.join(output_dir, f"MANDOROS_Phase3_Thesis_Navigator_{safe_topic}.xlsx")

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        # Styles
        navy_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cyan_accent_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10, color="000000")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        def style_worksheet(ws, df, sheet_title, is_matrix=False):
            ws.title = sheet_title
            # Title banner
            ws.merge_cells("A1:G1")
            banner_cell = ws["A1"]
            banner_cell.value = f"MANDOROS PHASE 3 — {sheet_title.upper().replace('_', ' ')}"
            banner_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
            banner_cell.fill = navy_header_fill
            banner_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 28

            # Table Header
            headers = list(df.columns)
            ws.append([])  # Row 2 is blank spacer
            ws.append(headers)  # Row 3
            header_row_idx = 3
            ws.row_dimensions[header_row_idx].height = 24

            for col_num in range(1, len(headers) + 1):
                c = ws.cell(row=header_row_idx, column=col_num)
                c.fill = cyan_accent_fill
                c.font = white_bold
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.border = thin_border

            # Data rows
            for r_idx, row in enumerate(df.itertuples(index=False), start=4):
                ws.row_dimensions[r_idx].height = 20
                for c_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    cell.value = str(val) if val is not None else ""
                    cell.font = regular_font
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

            # Auto fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(45, max(max_len + 3, 14))

        # ---------------------------------------------------------------------
        # Sheet 1: Evidence_Matrix
        # ---------------------------------------------------------------------
        matrix_rows = []
        for s in dossier.evidence_matrix.studies:
            matrix_rows.append({
                "Study ID": s.study_id,
                "Trial Title": s.title,
                "Condition Studied": s.condition,
                "Sample Size (N)": s.sample_size,
                "Treatment Arms": ", ".join(s.interventions),
                "Control / Comparator": ", ".join(s.comparators),
                "Primary Endpoints": ", ".join(s.primary_outcomes),
                "Secondary Endpoints": ", ".join(s.secondary_outcomes[:2]),
                "Follow-up (Wks)": s.follow_up_weeks,
                "Study Design": s.study_design,
                "Source File": s.source_file
            })
        df_matrix = pd.DataFrame(matrix_rows) if matrix_rows else pd.DataFrame([{"Notice": "No studies loaded"}])
        ws1 = wb.create_sheet()
        style_worksheet(ws1, df_matrix, "1_Evidence_Matrix")

        # ---------------------------------------------------------------------
        # Sheet 2: Observed_Research_Gaps
        # ---------------------------------------------------------------------
        gap_rows = []
        for g in dossier.evidence_matrix.observed_gaps:
            gap_rows.append({
                "Gap ID": g.gap_id,
                "Category": g.category,
                "Gap Title": g.title,
                "Description": g.description,
                "Observed In Studies": ", ".join(g.observed_in_studies),
                "Evidence Rationale": g.evidence_rationale,
                "Confidence Score": f"{g.confidence_score*100:.0f}%",
                "Verification Status": g.verification_status
            })
        df_gaps = pd.DataFrame(gap_rows) if gap_rows else pd.DataFrame([{"Notice": "No research gaps detected"}])
        ws2 = wb.create_sheet()
        style_worksheet(ws2, df_gaps, "2_Research_Gaps")

        # ---------------------------------------------------------------------
        # Sheet 3: Literature_Reality_Check
        # ---------------------------------------------------------------------
        rc_rows = []
        for rc in dossier.reality_check_reports:
            for m in rc.literature_matches + rc.registered_trials:
                rc_rows.append({
                    "Gap ID": rc.gap_id,
                    "Reality Check Status": rc.status,
                    "Database Source": m.source,
                    "Record Identifier": m.record_id,
                    "Publication / Trial Title": m.title,
                    "Year / Status": f"{m.year} ({m.overall_status})" if m.year else str(m.overall_status),
                    "Authors / Sponsor": m.authors_or_sponsor,
                    "Relevance Tier": m.relevance_tier,
                    "URL": m.url
                })
        df_rc = pd.DataFrame(rc_rows) if rc_rows else pd.DataFrame([{"Notice": "No external reality checks executed"}])
        ws3 = wb.create_sheet()
        style_worksheet(ws3, df_rc, "3_Reality_Check")

        # ---------------------------------------------------------------------
        # Sheet 4: Methodological_Appraisal
        # ---------------------------------------------------------------------
        rob_rows = []
        for a in dossier.methodological_assessments:
            for domain_name, details in a.domain_scores.items():
                rob_rows.append({
                    "Study ID": a.study_id,
                    "Framework": a.framework,
                    "Overall Quality Rating": a.overall_rating,
                    "Domain": domain_name,
                    "Domain Rating": details.get("rating", ""),
                    "Scientific Justification": details.get("justification", ""),
                    "Source Evidence Quote": details.get("evidence_quote", ""),
                    "Reviewer Overridden": "Yes" if a.reviewer_overridden else "No",
                    "Reviewer Notes": a.reviewer_notes or "N/A"
                })
        df_rob = pd.DataFrame(rob_rows) if rob_rows else pd.DataFrame([{"Notice": "No quality appraisals recorded"}])
        ws4 = wb.create_sheet()
        style_worksheet(ws4, df_rob, "4_Quality_Assessment_RoB")

        # ---------------------------------------------------------------------
        # Sheet 5: Thesis_Proposal_Blueprint
        # ---------------------------------------------------------------------
        tp_rows = []
        for p in dossier.thesis_proposals:
            for idx, title in enumerate(p.proposed_titles, start=1):
                tp_rows.append({
                    "Gap ID": p.gap_id,
                    "Title Option": f"Option {idx}",
                    "Proposed Academic Title": title,
                    "Recommended Design": p.recommended_study_design,
                    "Primary Objective": p.primary_objective,
                    "Population (P)": p.research_question_pico.get("Population", ""),
                    "Intervention (I)": p.research_question_pico.get("Intervention", ""),
                    "Comparator (C)": p.research_question_pico.get("Comparator", ""),
                    "Primary Outcome (O)": p.research_question_pico.get("Primary Outcome", ""),
                    "Setting (S)": p.research_question_pico.get("Setting", "")
                })
        df_tp = pd.DataFrame(tp_rows) if tp_rows else pd.DataFrame([{"Notice": "No thesis proposals formulated"}])
        ws5 = wb.create_sheet()
        style_worksheet(ws5, df_tp, "5_Thesis_Proposal")

        # ---------------------------------------------------------------------
        # Sheet 6: Sample_Size_Calculations
        # ---------------------------------------------------------------------
        ss_rows = []
        if dossier.sample_size_plan:
            sp = dossier.sample_size_plan
            ss_rows.append({
                "Study Design": sp.design_type,
                "Primary Endpoint": sp.primary_outcome_name,
                "Two-Sided Alpha (α)": sp.alpha,
                "Statistical Power (1-β)": f"{sp.power*100:.0f}%",
                "Allocation Ratio (r)": sp.allocation_ratio,
                "Anticipated Attrition": f"{sp.attrition_rate*100:.0f}%",
                "N Required Per Arm": sp.n_per_group,
                "Total Evaluable N": sp.n_total_evaluable,
                "Total Enrolled N (Adj)": sp.n_total_with_attrition,
                "Mathematical Formula": sp.mathematical_formula,
                "Assumptions & Rationale": sp.assumptions_rationale
            })
        df_ss = pd.DataFrame(ss_rows) if ss_rows else pd.DataFrame([{"Notice": "Sample size not calculated"}])
        ws6 = wb.create_sheet()
        style_worksheet(ws6, df_ss, "6_Sample_Size_Planning")

        # ---------------------------------------------------------------------
        # Sheet 7: Search_Strategy_PRISMA
        # ---------------------------------------------------------------------
        strat_rows = []
        if dossier.protocol_draft:
            pdraft = dossier.protocol_draft
            strat_rows.append({
                "Section": "Boolean Query Syntax",
                "Content": pdraft.boolean_query
            })
            strat_rows.append({
                "Section": "Target MeSH Terms",
                "Content": ", ".join(pdraft.mesh_terms_used)
            })
            for sec_name, sec_content in pdraft.sections.items():
                strat_rows.append({
                    "Section": sec_name,
                    "Content": sec_content
                })
        df_strat = pd.DataFrame(strat_rows) if strat_rows else pd.DataFrame([{"Notice": "Search strategy not drafted"}])
        ws7 = wb.create_sheet()
        style_worksheet(ws7, df_strat, "7_Protocol_and_Strategy")

        wb.save(excel_path)
        return excel_path

    @classmethod
    def export_markdown_report(
        cls,
        dossier: Phase3Dossier,
        output_dir: str = "outputs"
    ) -> str:
        """Generates a comprehensive Markdown research dossier."""
        os.makedirs(output_dir, exist_ok=True)
        safe_topic = dossier.target_topic.replace(" ", "_").replace("/", "_")[:30]
        md_path = os.path.join(output_dir, f"MANDOROS_Phase3_Dossier_{safe_topic}.md")

        content = []
        content.append(f"# 🧭 MANDOROS Phase 3: Research Gap & Thesis Navigator Dossier")
        content.append(f"**Target Topic:** {dossier.target_topic}  ")
        content.append(f"**Generated:** {dossier.generated_at}  ")
        content.append(f"**Evidence Base:** {len(dossier.evidence_matrix.studies)} Clinical Studies (N={dossier.evidence_matrix.total_patients_pooled} Total Analyzed Participants)  ")
        content.append("\n---\n")

        # Pillar I: Evidence Matrix & Gaps
        content.append("## Pillar I: Evidence Mapping & Observed Research Gaps")
        content.append(f"Across the analyzed studies, **{len(dossier.evidence_matrix.observed_gaps)} potential research gaps** were detected:\n")
        for g in dossier.evidence_matrix.observed_gaps:
            content.append(f"### 📍 [{g.gap_id}] {g.title}")
            content.append(f"- **Category:** `{g.category}`")
            content.append(f"- **Description:** {g.description}")
            content.append(f"- **Observed in Studies:** {', '.join(g.observed_in_studies)}")
            content.append(f"- **Scientific Rationale:** {g.evidence_rationale}")
            content.append(f"- **Confidence Score:** {g.confidence_score*100:.0f}%\n")

        # Pillar II: Reality Check
        content.append("## Pillar II: Global Scholarly & Clinical Trial Reality Check")
        for rc in dossier.reality_check_reports:
            content.append(f"### Reality Check for {rc.gap_id}")
            content.append(f"- **Search Query:** `{rc.search_query_used}`")
            content.append(f"- **Status:** **{rc.status}** ({rc.status_badge})")
            content.append(f"- **Summary Verdict:** {rc.summary_verdict}")
            content.append(f"- **External Evidence Matches:** {len(rc.literature_matches)} publications, {len(rc.registered_trials)} registered trials.\n")
            if rc.registered_trials:
                content.append("**Registered Clinical Trials Found:**")
                for t in rc.registered_trials[:3]:
                    content.append(f"- [{t.record_id}]({t.url}): *{t.title}* ({t.overall_status})")
            if rc.literature_matches:
                content.append("\n**Scholarly Literature Found:**")
                for m in rc.literature_matches[:3]:
                    content.append(f"- [{m.source}] *{m.title}* ({m.year or 'N/A'}) — {m.authors_or_sponsor}")
            content.append("\n")

        # Pillar III: Methodological Quality (RoB 2.0 / CONSORT)
        content.append("## Pillar III: Methodological Quality & Risk of Bias Appraisals")
        for a in dossier.methodological_assessments:
            content.append(f"### {a.study_id} — {a.framework}")
            content.append(f"- **Overall Rating:** **{a.overall_rating}** (CONSORT Reporting Score: {a.consort_reporting_score}%)")
            content.append(f"- **GRADE Certainty Level:** {a.grade_certainty}")
            content.append("| Domain | Rating | Scientific Justification |")
            content.append("| :--- | :--- | :--- |")
            for dom, det in a.domain_scores.items():
                content.append(f"| {dom} | **{det.get('rating')}** | {det.get('justification')} |")
            content.append("\n")

        # Pillar IV: Thesis Proposal Blueprint
        content.append("## Pillar IV: Thesis Proposal & Research Question Blueprint")
        for p in dossier.thesis_proposals:
            content.append(f"### Blueprint for {p.gap_id}")
            content.append("**Proposed Academic Titles:**")
            for idx, t in enumerate(p.proposed_titles, start=1):
                content.append(f"{idx}. *{t}*")
            content.append(f"\n- **Primary Objective:** {p.primary_objective}")
            content.append(f"- **Recommended Study Design:** {p.recommended_study_design}")
            content.append(f"- **Design Justification:** {p.design_justification}\n")
            content.append("#### PICO / PICOTS Framework:")
            for k, v in p.research_question_pico.items():
                content.append(f"- **{k}:** {v}")
            content.append("\n#### FINER Feasibility Evaluation:")
            for crit, det in p.finer_criteria.items():
                content.append(f"- **{crit}:** `{det.get('score')}` — {det.get('details')}")
            content.append("\n#### Claims Requiring Supervisor & IRB Approval:")
            for c in p.supervisor_verification_claims:
                content.append(f"- ⚠️ {c}")
            content.append("\n")

        # Pillar V: Statistical Planning & Search Strategy
        if dossier.sample_size_plan:
            sp = dossier.sample_size_plan
            content.append("## Pillar V: Biostatistical Sample Size Planning")
            content.append(f"- **Design:** `{sp.design_type}`")
            content.append(f"- **Primary Outcome:** {sp.primary_outcome_name}")
            content.append(f"- **Two-Sided Alpha (α):** {sp.alpha} | **Power (1-β):** {sp.power*100:.0f}%")
            content.append(f"- **Required Sample Size Per Arm (n):** **{sp.n_per_group}**")
            content.append(f"- **Total Evaluable Sample Size (N_eval):** **{sp.n_total_evaluable}**")
            content.append(f"- **Total Enrollment with {sp.attrition_rate*100:.0f}% Attrition (N_total):** **{sp.n_total_with_attrition}**")
            content.append(f"```text\n{sp.mathematical_formula}\n```")
            content.append(f"**Rationale:** {sp.assumptions_rationale}\n")

        if dossier.protocol_draft:
            pdraft = dossier.protocol_draft
            content.append("## Formal Protocol Draft (PRISMA-P / SPIRIT Outline)")
            content.append(f"**Framework:** {pdraft.framework}\n")
            content.append("### Reproducible PubMed MeSH Query:")
            content.append(f"```text\n{pdraft.boolean_query}\n```\n")
            for sec_name, sec_content in pdraft.sections.items():
                content.append(f"### {sec_name}")
                content.append(f"{sec_content}\n")

        report_text = "\n".join(content)
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(report_text)

        return md_path
