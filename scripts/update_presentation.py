import os
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = Path(__file__).resolve().parent.parent
PPTX_PATHS = [
    BASE_DIR / "MANDOROS_Global_Academic_Presentation.pptx",
    BASE_DIR / "mandoros_dossier" / "MANDOROS_Global_Academic_Presentation.pptx"
]

def add_architecture_slide(prs: Presentation):
    # Blank slide layout (usually layout index 6 in pptx default)
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Background color fill (Dark Slate: #0B1120)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(11, 17, 32)

    # 1. Category Tracker (Small top cyan text)
    top_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(0.4))
    tf_top = top_box.text_frame
    tf_top.word_wrap = True
    p_top = tf_top.paragraphs[0]
    p_top.text = "CLINICAL AI GOVERNANCE & ARCHITECTURE"
    p_top.font.name = "Arial"
    p_top.font.size = Pt(10)
    p_top.font.bold = True
    p_top.font.color.rgb = RGBColor(0, 245, 212) # Cyan

    # Top right Keynote badge
    keynote_box = slide.shapes.add_textbox(Inches(9.5), Inches(0.5), Inches(3), Inches(0.4))
    tf_k = keynote_box.text_frame
    p_k = tf_k.paragraphs[0]
    p_k.text = "MANDOROS PROTOCOL v2.0"
    p_k.alignment = PP_ALIGN.RIGHT
    p_k.font.name = "Arial"
    p_k.font.size = Pt(10)
    p_k.font.color.rgb = RGBColor(148, 163, 184)

    # 2. Slide Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = "Two-Phase Architecture & 10/10 Programmatic Quality Gate"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(248, 250, 252)

    # Subtitle
    p_sub = tf_title.add_paragraph()
    p_sub.text = "Enforcing Supreme Master Protocols on top of Dynamic Clinical Skills for Zero-Hallucination Evidence Synthesis."
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)

    # 3. Three Radiant Cards
    cards_data = [
        {
            "tag": "TIER 1: SUPREME LAW",
            "tag_color": RGBColor(239, 68, 68), # Coral/Red
            "title": "Master Protocols (Non-Negotiable)",
            "bullets": [
                "Phase 1 Extraction: 35 mandatory rules, zero duplication, PICO alignment, and CONSORT cell-level data normalization.",
                "Phase 2 Meta-Analysis: 21 verified worksheets, Cochrane RoB 2.0 (5 domains), and GRADE certainty framework.",
                "Absolute Authority: Tier 1 protocols override general skills or external LLM defaults 100% without exception.",
                "Zero Assumption Policy: Original research PDF is the supreme source of truth; zero extrapolation."
            ]
        },
        {
            "tag": "TIER 2: TECHNICAL MUSCLE",
            "tag_color": RGBColor(56, 189, 248), # Sky Blue
            "title": "Clinical Skills Library (Execution)",
            "bullets": [
                "Extraction Skills (Phase 1): 'pdf' (spatial table extraction), 'clinical-reports' (AE taxonomy), 'clinicaltrials-database'.",
                "Biostatistical Skills (Phase 2): 'statistical-analysis' (APA reporting), 'statsmodels' (mixed models, meta-regression).",
                "Decision Support & Review: 'clinical-decision-support' (evidence profiles) and 'peer-review' (CONSORT compliance).",
                "Dynamic runtime injection into system prompts for maximum mathematical rigor."
            ]
        },
        {
            "tag": "QUALITY GATE: 10/10",
            "tag_color": RGBColor(16, 185, 129), # Emerald Green
            "title": "Deterministic Validation & Retry Loop",
            "bullets": [
                "Hard Programmatic Quality Gate: Validates output against strict JSON schemas (phase1_schema & phase2_schema).",
                "Automated Self-Correction Loop: Feeds schema validation errors back to LLM for instant automated self-repair.",
                "Mathematical Sanity Verification: Enforces I² in [0, 100], τ² >= 0, CI lower <= CI upper, and positive ratio scales.",
                "Local === Render Cloud Parity: Path-agnostic relative routing across both Windows and Render Linux."
            ]
        }
    ]

    card_width = Inches(3.7)
    card_height = Inches(4.7)
    card_top = Inches(1.8)
    spacing = Inches(0.35)

    for i, c in enumerate(cards_data):
        left = Inches(0.8) + i * (card_width + spacing)
        
        # Background rectangle
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, card_top, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
        shape.line.color.rgb = c["tag_color"]
        shape.line.width = Pt(1.5)

        # Content inside shape
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)

        # Tag
        p_tag = tf.paragraphs[0]
        p_tag.text = c["tag"]
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = c["tag_color"]

        # Title
        p_t = tf.add_paragraph()
        p_t.text = c["title"]
        p_t.font.name = "Arial"
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(248, 250, 252)
        p_t.space_after = Pt(12)

        # Bullets
        for b in c["bullets"]:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.name = "Arial"
            p_b.font.size = Pt(10.5)
            p_b.font.color.rgb = RGBColor(203, 213, 225)
            p_b.space_after = Pt(8)

    # 4. Footer
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(8), Inches(0.3))
    tf_f = footer_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "MANDOROS • Smarter Evidence. Better Decisions. • Academic & Render Cloud Production Standard"
    p_f.font.name = "Arial"
    p_f.font.size = Pt(9)
    p_f.font.color.rgb = RGBColor(148, 163, 184)

    page_box = slide.shapes.add_textbox(Inches(10.5), Inches(6.8), Inches(2), Inches(0.3))
    tf_p = page_box.text_frame
    p_p = tf_p.paragraphs[0]
    p_p.text = f"SLIDE {len(prs.slides):02d} / {len(prs.slides):02d}"
    p_p.alignment = PP_ALIGN.RIGHT
    p_p.font.name = "Arial"
    p_p.font.size = Pt(9)
    p_p.font.bold = True
    p_p.font.color.rgb = RGBColor(56, 189, 248)

def update_all_presentations():
    for p_path in PPTX_PATHS:
        if p_path.exists():
            prs = Presentation(str(p_path))
            add_architecture_slide(prs)
            prs.save(str(p_path))
            print(f"[+] Successfully added Architecture & Quality Gate slide to: {p_path} (Now {len(prs.slides)} slides)")

if __name__ == "__main__":
    update_all_presentations()
