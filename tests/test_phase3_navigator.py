"""
Unit & Integration Test Suite for Phase 3: Research Gap & Thesis Navigator (Global Standards Edition).
Tests the strict deterministic backend components in phase3_navigator:
- SchemaIngestor
- GapTaxonomyExtractor (4-Type Taxonomy)
- GlobalRealityChecker (3-Color Traffic-Light)
- EGMRobvisPlotter (EGM & Cochrane robvis)
- SampleSizeAndRiskEngine (presize & Wood-McNair risk)
- ThesisProtocolBuilder (FINER, GRADE, PROSPERO, SPIRIT)
- Phase3DossierExporter (6-sheet color-coded Excel & HTML)
- ui_tab3 pipeline integration
"""

import os
import sys
from pathlib import Path
import pytest
import pandas as pd
import openpyxl
import plotly.graph_objects as go

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from phase3_navigator.schema_ingestor import SchemaIngestor, IngestedTrialData
from phase3_navigator.gap_taxonomy_extractor import (
    GapTaxonomyExtractor,
    ScientificGap,
    GapCategory,
)
from phase3_navigator.global_reality_checker import (
    GlobalRealityChecker,
    RealityCheckVerdict,
    TrafficLightColor,
)
from phase3_navigator.egm_robvis_plotter import EGMRobvisPlotter
from phase3_navigator.sample_size_and_risk import (
    SampleSizeAndRiskEngine,
    SampleSizeResult,
    SampleSizeSensitivityTable,
    ProtocolRiskReport,
)
from phase3_navigator.thesis_protocol_builder import (
    ThesisProtocolBuilder,
    ThesisBlueprint,
    GRADESummaryOfFindings,
)
from phase3_navigator.excel_html_exporter import Phase3DossierExporter
from phase3_navigator.ui_tab3 import run_phase3_pipeline


@pytest.fixture
def sample_ingested_trial():
    return IngestedTrialData(
        study_title="A 52-Week Randomized Trial of Tirzepatide Versus Placebo in Type 2 Diabetes",
        condition="Type 2 Diabetes Mellitus",
        drug_name="Tirzepatide",
        comparator_name="Placebo",
        total_sample_size=475,
        arms=[
            {"name": "Tirzepatide 15mg", "sample_size": 238},
            {"name": "Matching Placebo", "sample_size": 237},
        ],
        outcomes=[
            {"name": "Change in HbA1c (%)", "is_primary": True, "timepoint": "Week 52", "effect_size": -2.07},
            {"name": "Body Weight Change (kg)", "is_primary": False, "timepoint": "Week 52", "effect_size": -9.5},
        ],
        inclusion_criteria=["Age 18-75 years", "HbA1c 7.5-10.5%", "BMI >= 25 kg/m2"],
        exclusion_criteria=["eGFR < 30 mL/min/1.73m2", "Age >= 75 years", "History of pancreatitis"],
        extracted_tables={"Baseline": pd.DataFrame({"Metric": ["Mean Age", "Female %"], "Value": [57.2, 48.5]})},
        provenance_map={"HbA1c": "Efficacy Outcomes [Row 12]", "Exclusion": "Eligibility Criteria [Row 24]"},
    )


# -----------------------------------------------------------------------------
# 1. Gap Taxonomy Extractor Tests
# -----------------------------------------------------------------------------
def test_gap_taxonomy_extraction_4_categories(sample_ingested_trial):
    """Verifies that gaps are classified into all 4 scientific categories."""
    gaps = GapTaxonomyExtractor.extract_gaps(sample_ingested_trial)
    assert len(gaps) >= 4

    categories = {g.category for g in gaps}
    assert GapCategory.POP_GAP in categories
    assert GapCategory.INT_COMP_GAP in categories
    assert GapCategory.OUTCOME_GAP in categories
    assert GapCategory.METHOD_ROB_GRADE_GAP in categories

    for g in gaps:
        assert g.gap_id.startswith("GAP_") or g.gap_id.startswith("GAP-")
        assert g.badge_label != ""
        assert g.badge_arabic != ""
        assert g.source_sheet != ""
        assert g.row_coord != ""
        assert g.exact_extracted_value_or_quote != ""
        assert g.target_query_keyword != ""


# -----------------------------------------------------------------------------
# 2. Global Reality Checker Tests
# -----------------------------------------------------------------------------
def test_global_reality_checker_offline_resilience(sample_ingested_trial):
    """Verifies reality check produces valid verdict even when network is skipped/offline."""
    gaps = GapTaxonomyExtractor.extract_gaps(sample_ingested_trial)
    gap = gaps[0]

    verdict = GlobalRealityChecker.verify_gap(
        gap=gap,
        condition=sample_ingested_trial.condition,
        drug_name=sample_ingested_trial.drug_name,
        skip_network=True,
    )
    assert isinstance(verdict, RealityCheckVerdict)
    assert verdict.gap_id == gap.gap_id
    assert verdict.traffic_light_color in [TrafficLightColor.GREEN, TrafficLightColor.YELLOW, TrafficLightColor.RED]
    assert verdict.badge_label != ""
    assert verdict.actionable_recommendation != ""


# -----------------------------------------------------------------------------
# 3. Evidence & Gap Map & robvis Plotter Tests
# -----------------------------------------------------------------------------
def test_egm_bubble_matrix_generation(sample_ingested_trial):
    """Verifies Plotly EGM matrix generation follows Campbell/EviAtlas standard."""
    gaps = GapTaxonomyExtractor.extract_gaps(sample_ingested_trial)
    fig = EGMRobvisPlotter.build_egm_bubble_matrix(gaps, sample_ingested_trial)
    assert isinstance(fig, go.Figure)
    assert len(fig.data) >= 1  # Contains bubble scatter + legend traces


def test_cochrane_rob2_traffic_light_plot(sample_ingested_trial):
    """Verifies McGuinness & Higgins 2021 Cochrane robvis traffic-light chart."""
    fig = EGMRobvisPlotter.build_rob2_traffic_light_plot(sample_ingested_trial)
    assert isinstance(fig, go.Figure)
    assert "Cochrane" in fig.layout.title.text


def test_cochrane_rob2_summary_bar(sample_ingested_trial):
    """Verifies McGuinness & Higgins 2021 Cochrane robvis summary bar chart."""
    fig = EGMRobvisPlotter.build_rob2_summary_bar(sample_ingested_trial)
    assert isinstance(fig, go.Figure)
    assert fig.layout.barmode == "stack"


# -----------------------------------------------------------------------------
# 4. Biostatistical Sample Size & Protocol Risk Tests
# -----------------------------------------------------------------------------
def test_sample_size_continuous_superiority():
    """Delta = 5.0, SD = 10.0 (Cohen's d = 0.50), alpha = 0.05, power = 0.80 -> n ~ 63-64."""
    res = SampleSizeAndRiskEngine.calculate_continuous_superiority(
        delta=5.0,
        sigma=10.0,
        alpha=0.05,
        power=0.80,
        attrition_rate=0.15,
    )
    assert 62 <= res.n_per_arm_evaluable <= 65
    assert res.n_total_evaluable == res.n_per_arm_evaluable * 2
    assert res.n_per_arm_enrolled > res.n_per_arm_evaluable


def test_sample_size_binary_superiority():
    """p1 = 0.45, p2 = 0.25, alpha = 0.05, power = 0.80 -> n ~ 80-100."""
    res = SampleSizeAndRiskEngine.calculate_binary_superiority(
        p_treatment=0.45,
        p_control=0.25,
        alpha=0.05,
        power=0.80,
    )
    assert 80 <= res.n_per_arm_evaluable <= 100


def test_sample_size_sensitivity_matrix():
    """Verifies multi-scenario sensitivity grid generation across power and alpha."""
    base = SampleSizeAndRiskEngine.calculate_continuous_superiority(delta=5.0, sigma=10.0)
    matrix = SampleSizeAndRiskEngine.build_sensitivity_matrix(base, delta=5.0, sigma=10.0)
    assert len(matrix.scenarios) == 12  # 3 effect sizes * 2 alphas * 2 powers
    assert matrix.dataframe is not None
    assert len(matrix.dataframe) == 12


def test_protocol_risk_auditor(sample_ingested_trial):
    """Verifies Wood & McNair 2023 5-factor protocol risk audit (0-100)."""
    report = SampleSizeAndRiskEngine.audit_protocol_risk(sample_ingested_trial)
    assert 0.0 <= report.total_score <= 100.0
    assert len(report.domains) == 5
    assert report.risk_category in [
        "Low Risk of Uninformative Trial",
        "Moderate Risk (Methodological Caution)",
        "High Risk of Trial Uninformativeness"
    ]


# -----------------------------------------------------------------------------
# 5. Thesis Protocol Builder Tests (GRADE, PROSPERO, SPIRIT)
# -----------------------------------------------------------------------------
def test_thesis_protocol_builder(sample_ingested_trial):
    """Verifies FINER criteria, GRADE table, and PROSPERO/SPIRIT drafts."""
    gaps = GapTaxonomyExtractor.extract_gaps(sample_ingested_trial)
    blueprint = ThesisProtocolBuilder.build_blueprint(gaps[0], sample_ingested_trial)

    assert len(blueprint.proposed_academic_titles) == 3
    assert "Population (P)" in blueprint.pico_framework
    assert len(blueprint.finer_criteria) == 5
    assert blueprint.finer_overall_score > 0
    assert len(blueprint.grade_table.rows) >= 3
    assert "MeSH" in blueprint.pubmed_boolean_query
    assert "1. Review Title" in blueprint.prospero_protocol_draft
    assert "Item 1: Title" in blueprint.spirit_protocol_synopsis


# -----------------------------------------------------------------------------
# 6. Exporter Tests (6-Sheet Color-Coded Excel & Standalone HTML)
# -----------------------------------------------------------------------------
def test_phase3_excel_and_html_export(sample_ingested_trial, tmp_path):
    """Verifies 6-sheet color-coded Excel and standalone HTML generation."""
    gaps = GapTaxonomyExtractor.extract_gaps(sample_ingested_trial)
    base = SampleSizeAndRiskEngine.calculate_continuous_superiority(delta=5.0, sigma=10.0)
    matrix = SampleSizeAndRiskEngine.build_sensitivity_matrix(base, delta=5.0, sigma=10.0)
    risk = SampleSizeAndRiskEngine.audit_protocol_risk(sample_ingested_trial)
    blueprint = ThesisProtocolBuilder.build_blueprint(gaps[0], sample_ingested_trial)

    excel_file = str(tmp_path / "test_phase3_export.xlsx")
    html_file = str(tmp_path / "test_phase3_export.html")

    Phase3DossierExporter.export_excel(
        output_path=excel_file,
        gaps=gaps,
        sensitivity_table=matrix,
        risk_report=risk,
        blueprint=blueprint,
        trial_data=sample_ingested_trial,
    )
    assert os.path.exists(excel_file)

    # Check 6 sheets in Excel
    wb = openpyxl.load_workbook(excel_file)
    expected_sheets = [
        "Executive Gap Matrix",
        "Evidence & Gap Map Data",
        "Sample Size & Power Models",
        "Protocol Risk Audit",
        "GRADE Summary of Findings",
        "Search Strategies & Blueprints",
    ]
    assert len(wb.sheetnames) == 6
    for s in expected_sheets:
        assert s in wb.sheetnames

    # Check HTML export
    Phase3DossierExporter.export_html_dossier(
        output_path=html_file,
        gaps=gaps,
        sensitivity_table=matrix,
        risk_report=risk,
        blueprint=blueprint,
        trial_data=sample_ingested_trial,
    )
    assert os.path.exists(html_file)
    with open(html_file, "r", encoding="utf-8") as f:
        html_text = f.read()
    assert "MANDOROS Phase 3" in html_text
    assert "Evidence &amp; Gap Map" in html_text


# -----------------------------------------------------------------------------
# 7. UI Pipeline Enforcement: No File = No Fallback / No Synthetic Data
# -----------------------------------------------------------------------------
def test_ui_pipeline_enforces_no_empty_upload():
    """Verifies that clicking run without files does not use fake benchmark data."""
    outputs = run_phase3_pipeline(
        uploaded_files=None,
        selected_gap_choice="Gap 1",
        design_choice="Continuous",
        alpha=0.05,
        power=0.80,
        attrition=0.15,
        mean_diff=5.0,
        pooled_sd=10.0,
        p1_rate=0.45,
        p2_rate=0.25,
        ni_margin=2.5,
    )
    status_html, egm_fig, rob_fig, gaps_df, rc_html, sens_df, risk_html, grade_df, bp_html, xl_path, html_path = outputs

    assert "Missing Uploaded Trial Files" in status_html
    assert xl_path is None
    assert html_path is None
    assert egm_fig is None
