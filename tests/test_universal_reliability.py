"""
Comprehensive Verification Suite for MANDOROS Universal Reliability Engineering.
Tests:
1. Requirements Registry (Category A deterministic and Category B scientific rules).
2. Quality Gates A through H.
3. Gate E Workbook Reopening & Post-Export Cell Reconciliation.
4. Single-Study Protection (k=1 safeguard).
5. 7-State Missingness Taxonomy (Zero-Conflation of extraction failure vs unmeasured).
6. Universal Unseen Trial Extraction & Harmonization.
"""

import os
import sys
from pathlib import Path
import pytest
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.models import (
    StudyMetadata,
    TreatmentArm,
    EligibilityCriterion,
    OutcomeMeasurement,
    ProvenanceRecord,
    ValidationIssue,
    ExtractedTrialData,
)
from src.requirements_registry import (
    MissingDataStatus,
    RequirementCategory,
    RequirementsRegistry,
    registry,
)
from src.quality_gates import QualityGates, GateResult
from src.exporter import ExcelTrialExporter


# -------------------------------------------------------------------------
# FIXTURES
# -------------------------------------------------------------------------
@pytest.fixture
def sample_trial_data() -> ExtractedTrialData:
    study = StudyMetadata(
        study_id="RCT_NOVEL_DRUG_2026",
        title="Efficacy and Safety of Novel Compound X in Cardiometabolic Disease: A Randomized Double-Blind Trial",
        authors="Al-Mansoor et al.",
        publication_year=2026,
        journal="New England Journal of Medicine",
        doi="10.1056/NEJMoa2026001",
        study_design="Randomized, double-blind, placebo-controlled trial",
        study_phase="Phase III",
        total_sample_size=300,
        condition_studied="Type 2 Diabetes and Obesity"
    )

    arms = [
        TreatmentArm(
            arm_id="ARM_001",
            study_id=study.study_id,
            arm_name="Compound X 50mg",
            intervention_type="experimental",
            drug_name="Compound X",
            dose_value=50.0,
            dose_unit="mg",
            frequency_number=1.0,
            frequency_unit="daily",
            sample_size_allocated=150,
            sample_size_analyzed=145,
            source_page=2
        ),
        TreatmentArm(
            arm_id="ARM_002",
            study_id=study.study_id,
            arm_name="Placebo",
            intervention_type="placebo",
            drug_name="Placebo",
            dose_value=None,
            dose_unit=None,
            frequency_number=1.0,
            frequency_unit="daily",
            sample_size_allocated=150,
            sample_size_analyzed=142,
            source_page=2
        )
    ]

    eligibility = [
        EligibilityCriterion(
            criterion_id="ELIG_001",
            study_id=study.study_id,
            criterion_type="inclusion",
            domain="demographic",
            criterion_name="age",
            min_operator=">=",
            min_value=18.0,
            max_operator="<=",
            max_value=75.0,
            unit="years",
            raw_source_text="Age 18 to 75 years",
            source_page=3
        ),
        EligibilityCriterion(
            criterion_id="ELIG_002",
            study_id=study.study_id,
            criterion_type="inclusion",
            domain="glycemic",
            criterion_name="hba1c",
            min_operator=">=",
            min_value=7.0,
            max_operator="<=",
            max_value=10.5,
            unit="%",
            raw_source_text="HbA1c between 7.0% and 10.5%",
            source_page=3
        )
    ]

    outcomes = [
        # Baseline HbA1c
        OutcomeMeasurement(
            outcome_id="OUT_001",
            study_id=study.study_id,
            arm_id="ARM_001",
            domain="glycemic",
            parameter_name="HbA1c",
            statistical_type="continuous",
            timepoint_value="0",
            timepoint_unit="weeks",
            timepoint="Baseline",
            sample_size=145,
            mean=8.2,
            sd=0.9,
            unit="%",
            is_baseline=True,
            source_page=4,
            raw_source_text="Baseline HbA1c: 8.2 ± 0.9%"
        ),
        OutcomeMeasurement(
            outcome_id="OUT_002",
            study_id=study.study_id,
            arm_id="ARM_002",
            domain="glycemic",
            parameter_name="HbA1c",
            statistical_type="continuous",
            timepoint_value="0",
            timepoint_unit="weeks",
            timepoint="Baseline",
            sample_size=142,
            mean=8.3,
            sd=0.8,
            unit="%",
            is_baseline=True,
            source_page=4,
            raw_source_text="Baseline HbA1c: 8.3 ± 0.8%"
        ),
        # Follow-up HbA1c Change
        OutcomeMeasurement(
            outcome_id="OUT_003",
            study_id=study.study_id,
            arm_id="ARM_001",
            domain="glycemic",
            parameter_name="HbA1c",
            statistical_type="continuous",
            timepoint_value="24",
            timepoint_unit="weeks",
            timepoint="24 weeks",
            sample_size=145,
            change_value=-1.25,
            change_sd=0.45,
            unit="%",
            is_baseline=False,
            is_primary=True,
            effect_measure="Adjusted Mean Difference vs Placebo",
            effect_value=-0.85,
            ci_lower=-1.10,
            ci_upper=-0.60,
            p_value="<0.001",
            source_page=5,
            raw_source_text="Mean change in HbA1c was -1.25 ± 0.45% (Diff -0.85, P<0.001)"
        ),
        OutcomeMeasurement(
            outcome_id="OUT_004",
            study_id=study.study_id,
            arm_id="ARM_002",
            domain="glycemic",
            parameter_name="HbA1c",
            statistical_type="continuous",
            timepoint_value="24",
            timepoint_unit="weeks",
            timepoint="24 weeks",
            sample_size=142,
            change_value=-0.40,
            change_sd=0.50,
            unit="%",
            is_baseline=False,
            source_page=5,
            raw_source_text="Mean change in placebo was -0.40 ± 0.50%"
        ),
        # Safety Outcome: Hypoglycemia
        OutcomeMeasurement(
            outcome_id="OUT_005",
            study_id=study.study_id,
            arm_id="ARM_001",
            domain="safety_adverse_events",
            parameter_name="Hypoglycemia",
            statistical_type="dichotomous",
            timepoint_value="24",
            timepoint_unit="weeks",
            timepoint="24 weeks",
            sample_size=150,
            events=6,
            total=150,
            percentage=4.0,
            unit="%",
            is_baseline=False,
            source_page=6,
            raw_source_text="Hypoglycemia occurred in 6 patients (4.0%)"
        ),
        OutcomeMeasurement(
            outcome_id="OUT_006",
            study_id=study.study_id,
            arm_id="ARM_002",
            domain="safety_adverse_events",
            parameter_name="Hypoglycemia",
            statistical_type="dichotomous",
            timepoint_value="24",
            timepoint_unit="weeks",
            timepoint="24 weeks",
            sample_size=150,
            events=5,
            total=150,
            percentage=3.3,
            unit="%",
            is_baseline=False,
            source_page=6,
            raw_source_text="Hypoglycemia occurred in 5 patients (3.3%)"
        )
    ]

    provenance = [
        ProvenanceRecord(
            provenance_id="PROV_001",
            study_id=study.study_id,
            target_entity="Studies",
            target_id=study.study_id,
            field_name="title",
            source_page=1,
            source_location="Title Page",
            exact_source_text=study.title
        )
    ]

    validation_log = []

    return ExtractedTrialData(
        study=study,
        arms=arms,
        eligibility=eligibility,
        outcomes=outcomes,
        provenance=provenance,
        validation_log=validation_log
    )


# -------------------------------------------------------------------------
# 1. REQUIREMENTS REGISTRY TESTS
# -------------------------------------------------------------------------
def test_requirements_registry_categories():
    """Verifies that RequirementsRegistry categorizes rules into Category A & B."""
    cat_a = registry.get_category_a_rules()
    cat_b = registry.get_category_b_rules()

    assert len(cat_a) >= 5, "Category A deterministic rules must be registered"
    assert len(cat_b) >= 5, "Category B scientific interpretation rules must be registered"

    # Verify specific rule presence
    assert any(r.rule_id == "REQ_DET_001" for r in cat_a)
    assert any(r.rule_id == "REQ_SCI_001" for r in cat_b)


def test_missing_data_classification_taxonomy():
    """Verifies that missingness states are accurately separated without conflation."""
    # 1. Successfully extracted value
    s1 = registry.classify_missingness("hba1c", "HbA1c was 8.2%", 8.2, field_searched=True)
    assert s1 == MissingDataStatus.REPORTED_EXTRACTED

    # 2. Variable searched thoroughly but unmeasured in paper
    s2 = registry.classify_missingness("alt", "Table 1 lists demographic only", None, field_searched=True, extraction_attempted=True)
    assert s2 == MissingDataStatus.NOT_REPORTED_AFTER_SEARCH

    # 3. Extraction failed (unattempted or failed pipeline stage)
    s3 = registry.classify_missingness("bmi", "", None, field_searched=False, extraction_attempted=False)
    assert s3 == MissingDataStatus.EXTRACTION_FAILED

    # 4. Single-study meta-analysis variable (scientifically inapplicable)
    s4 = registry.classify_missingness("pooled_effect", "", None, study_design="Single study trial")
    assert s4 == MissingDataStatus.NOT_APPLICABLE


# -------------------------------------------------------------------------
# 2. QUALITY GATES A THROUGH H TESTS
# -------------------------------------------------------------------------
def test_gate_a_ingestion_pre_flight(tmp_path):
    """Gate A: Validates PDF header magic bytes and file size."""
    # Test valid PDF dummy
    valid_pdf = tmp_path / "test_trial.pdf"
    valid_pdf.write_bytes(b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n")

    res_valid = QualityGates.enforce_gate_a(str(valid_pdf))
    assert res_valid.passed is True
    assert res_valid.gate_id == "GATE_A"
    assert res_valid.metrics["file_size_bytes"] > 0
    assert "sha256" in res_valid.metrics

    # Test invalid file (not a PDF)
    invalid_file = tmp_path / "bad.pdf"
    invalid_file.write_bytes(b"INVALID_HEADER_DATA_NOT_PDF")
    res_bad = QualityGates.enforce_gate_a(str(invalid_file))
    assert res_bad.passed is False
    assert any("magic bytes" in e for e in res_bad.errors)


def test_gate_b_extraction_completeness(sample_trial_data):
    """Gate B: Validates presence of StudyMetadata, TreatmentArms, Outcomes, and Provenance."""
    res = QualityGates.enforce_gate_b(sample_trial_data)
    assert res.passed is True
    assert res.score == 100.0
    assert res.metrics["arms_count"] == 2
    assert res.metrics["outcomes_count"] == 6


def test_gate_c_normalization_and_units(sample_trial_data):
    """Gate C: Validates strict separation of value and unit, and percentage bounds."""
    res = QualityGates.enforce_gate_c(sample_trial_data)
    assert res.passed is True
    assert res.metrics["units_audited"] == 6

    # Test failure injection: unit inside numeric mean
    corrupted_data = sample_trial_data.model_copy(deep=True)
    corrupted_data.outcomes[0].mean = "8.2 %"  # String contaminated with unit
    res_corrupted = QualityGates.enforce_gate_c(corrupted_data)
    assert res_corrupted.passed is False
    assert any("has unit in numeric field" in e for e in res_corrupted.errors)


def test_gate_d_workbook_structure():
    """Gate D: Validates canonical sheets and ordering before saving."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Add required sheets
    wb.create_sheet("Master Excel Sheet")
    wb.create_sheet("Studies")
    wb.create_sheet("Arms")
    wb.create_sheet("Eligibility")
    wb.create_sheet("Outcomes")
    wb.create_sheet("Provenance")
    wb.create_sheet("Validation_Log")

    res = QualityGates.enforce_gate_d(wb, phase=1)
    assert res.passed is True
    assert res.metrics["sheet_count"] == 7

    # Remove Master Excel Sheet to trigger failure
    wb.remove(wb["Master Excel Sheet"])
    res_fail = QualityGates.enforce_gate_d(wb, phase=1)
    assert res_fail.passed is False


# -------------------------------------------------------------------------
# 3. GATE E: WORKBOOK REOPENING & POST-EXPORT RECONCILIATION TEST (CRITICAL)
# -------------------------------------------------------------------------
def test_gate_e_workbook_reopening_and_reconciliation(sample_trial_data, tmp_path):
    """
    Gate E: Exports trial data, saves to disk, reopens with openpyxl,
    and asserts 100% data reconciliation and zero cell corruption.
    """
    out_xlsx = str(tmp_path / "Universal_Reliability_Export.xlsx")
    exporter = ExcelTrialExporter(sample_trial_data)
    exported_path = exporter.export(out_xlsx)

    assert os.path.exists(exported_path)

    # Enforce Gate E explicitly
    res_gate_e = QualityGates.enforce_gate_e(exported_path, sample_trial_data)
    assert res_gate_e.passed is True, f"Gate E failed: {res_gate_e.errors}"
    assert res_gate_e.score >= 90.0
    assert res_gate_e.metrics["total_cells_audited"] > 50
    assert res_gate_e.metrics["corrupted_cells_count"] == 0
    assert "Master Excel Sheet" in res_gate_e.metrics["sheets_found"]
    assert "Arms" in res_gate_e.metrics["sheets_found"]
    assert "Outcomes" in res_gate_e.metrics["sheets_found"]

    # Reopen manually to assert specific cell values
    wb = openpyxl.load_workbook(exported_path, data_only=True)
    ws_master = wb["Master Excel Sheet"]
    assert ws_master.max_row >= 5, "Master Excel Sheet must have detailed content"

    # Verify no cell contains literal 'None' string
    for row in ws_master.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str):
                assert cell.strip().lower() not in ["none", "nan", "#value!", "#ref!"], f"Corrupted cell value found: {cell}"


# -------------------------------------------------------------------------
# 4. GATE F: SINGLE-STUDY (k=1) SYNTHESIS SAFEGUARD TEST
# -------------------------------------------------------------------------
def test_gate_f_single_study_safeguard():
    """Gate F: Asserts that k=1 prevents fake pooling and enforces NOT_APPLICABLE_SINGLE_STUDY."""
    # Single study with proper safeguard flag
    synth_single_ok = {
        "is_single_study": True,
        "meta_status": "NOT_APPLICABLE_SINGLE_STUDY",
        "notes": "Single study: pooling not applicable."
    }
    res_single = QualityGates.enforce_gate_f(study_count=1, synthesis_data=synth_single_ok)
    assert res_single.passed is True
    assert res_single.metrics["mode"] == "SINGLE_STUDY_DESCRIPTIVE"
    assert res_single.metrics["single_study_protected"] is True

    # Single study that erroneously claims a pooled effect
    synth_single_bad = {
        "pooled_effect": -0.85,
        "is_single_study": False,
        "meta_status": "POOLED"
    }
    res_bad = QualityGates.enforce_gate_f(study_count=1, synthesis_data=synth_single_bad)
    assert res_bad.passed is False
    assert any("Single study (k=1) claimed pooled meta-analysis" in e for e in res_bad.errors)

    # Multi-study (k >= 2) synthesis
    synth_multi = {
        "pooled_effect": -0.75,
        "i_squared": 25.4,
        "is_single_study": False
    }
    res_multi = QualityGates.enforce_gate_f(study_count=3, synthesis_data=synth_multi)
    assert res_multi.passed is True
    assert res_multi.metrics["mode"] == "MULTI_STUDY_META_ANALYSIS"
    assert res_multi.metrics["degrees_of_freedom"] == 2


# -------------------------------------------------------------------------
# 5. GATE G: PHASE 3 GAP & PREDICTIVE PRECONDITIONS TEST
# -------------------------------------------------------------------------
def test_gate_g_phase3_sample_size_preconditions():
    """Gate G: Validates alpha, power, and sample size output bounds."""
    # Valid input
    inputs_ok = {"alpha": 0.05, "power": 0.80}
    results_ok = {"total_sample_size": 250, "sample_size_per_arm": 125}
    res_ok = QualityGates.enforce_gate_g(inputs_ok, results_ok)
    assert res_ok.passed is True

    # Invalid alpha (alpha = 0.6)
    inputs_bad_alpha = {"alpha": 0.60, "power": 0.80}
    res_bad = QualityGates.enforce_gate_g(inputs_bad_alpha)
    assert res_bad.passed is False
    assert any("alpha" in e for e in res_bad.errors)


# -------------------------------------------------------------------------
# 6. GATE H: FINAL ARTIFACT INTEGRITY TEST
# -------------------------------------------------------------------------
def test_gate_h_final_artifact_readiness(tmp_path, sample_trial_data):
    """Gate H: Validates non-zero file size, zip/xlsx signature, and read permissions."""
    out_xlsx = str(tmp_path / "Final_Artifact_Test.xlsx")
    exporter = ExcelTrialExporter(sample_trial_data)
    exported_path = exporter.export(out_xlsx)

    res_h = QualityGates.enforce_gate_h(exported_path)
    assert res_h.passed is True
    assert res_h.metrics["file_size_bytes"] > 2048


# -------------------------------------------------------------------------
# 7. UNIVERSAL ORCHESTRATOR FULL AUDIT CERTIFICATE TEST
# -------------------------------------------------------------------------
def test_universal_pipeline_audit_certificate(tmp_path, sample_trial_data):
    """Verifies that QualityGates.audit_full_pipeline produces a certified audit."""
    out_xlsx = str(tmp_path / "Cert_Test.xlsx")
    exporter = ExcelTrialExporter(sample_trial_data)
    exported_path = exporter.export(out_xlsx)

    cert = QualityGates.audit_full_pipeline(
        data=sample_trial_data,
        xlsx_path=exported_path,
        k_studies=1
    )

    assert cert.overall_status in ["CERTIFIED", "CERTIFIED_WITH_WARNINGS"]
    assert cert.overall_score >= 85.0
    assert "GATE_B" in cert.gates
    assert "GATE_C" in cert.gates
    assert "GATE_E" in cert.gates
    assert "GATE_F" in cert.gates
    assert "GATE_H" in cert.gates
    assert cert.reconciliation_summary is not None
