import pytest
import os
import sys
import openpyxl

sys.path.insert(0, ".")

from src.extractor import RCTExtractor
from src.exporter import ExcelTrialExporter


@pytest.fixture(scope="module")
def rct4_extracted_data():
    pdf_path = "data/raw/empa vs placebo 2doses (RCT4).pdf"
    assert os.path.exists(pdf_path), f"Missing test PDF at {pdf_path}"
    ext = RCTExtractor(pdf_path)
    data = ext.extract_all(study_id_override="RCT4")
    return data


def test_study_metadata_and_phase(rct4_extracted_data):
    """Verifies Errors 2, 3: Study phase is strictly 'Not explicitly reported' and duration is 28 days."""
    study = rct4_extracted_data.study
    assert "Effect of empagliflozin monotherapy" in study.title
    assert study.doi == "10.1186/s12933-014-0169-9"
    assert study.journal == "Cardiovascular Diabetology"
    assert study.study_phase == "Not explicitly reported"
    assert study.duration_value == 28.0
    assert study.duration_unit == "days"
    assert study.total_sample_size == 60


def test_treatment_arms_and_sample_sizes(rct4_extracted_data):
    """Verifies Error 1: All 3 arms are retained with exact sample sizes (21, 20, 19)."""
    arms = rct4_extracted_data.arms
    assert len(arms) == 3
    arm_names = [a.arm_name for a in arms]
    assert "Placebo" in arm_names
    assert "Empagliflozin 10 mg" in arm_names
    assert "Empagliflozin 25 mg" in arm_names

    p_arm = next(a for a in arms if a.intervention_type == "placebo")
    assert p_arm.sample_size_allocated == 21

    e10_arm = next(a for a in arms if a.dose_value == 10.0)
    assert e10_arm.sample_size_allocated == 20

    e25_arm = next(a for a in arms if a.dose_value == 25.0)
    assert e25_arm.sample_size_allocated == 19


def test_table1_baselines_sd_dispersion(rct4_extracted_data):
    """Verifies Errors 4-7: Table 1 baselines have exact values and explicit SD dispersion."""
    outcomes = rct4_extracted_data.outcomes
    t1_map = {}
    for o in outcomes:
        if o.is_baseline and o.source_location == "Table 1":
            t1_map[(o.parameter_name, o.arm_id)] = o

    # BMI: 24.9 (2.8), 24.1 (3.7), 24.0 (3.2)
    assert t1_map[("BMI", "ARM_001")].mean == 24.9
    assert t1_map[("BMI", "ARM_001")].sd == 2.8
    assert t1_map[("BMI", "ARM_001")].dispersion_type == "SD"
    assert t1_map[("BMI", "ARM_002")].mean == 24.1
    assert t1_map[("BMI", "ARM_002")].sd == 3.7
    assert t1_map[("BMI", "ARM_003")].mean == 24.0
    assert t1_map[("BMI", "ARM_003")].sd == 3.2

    # PPG AUC1-4h: 682.8 (91.2), 680.4 (92.2), 658.1 (116.2)
    assert t1_map[("PPG AUC1-4h", "ARM_001")].mean == 682.8
    assert t1_map[("PPG AUC1-4h", "ARM_001")].sd == 91.2
    assert t1_map[("PPG AUC1-4h", "ARM_002")].mean == 680.4
    assert t1_map[("PPG AUC1-4h", "ARM_002")].sd == 92.2
    assert t1_map[("PPG AUC1-4h", "ARM_003")].mean == 658.1
    assert t1_map[("PPG AUC1-4h", "ARM_003")].sd == 116.2

    # MAGE: 91.4 (26.4), 94.1 (18.5), 89.1 (29.7)
    assert t1_map[("MAGE", "ARM_001")].mean == 91.4
    assert t1_map[("MAGE", "ARM_001")].sd == 26.4
    assert t1_map[("MAGE", "ARM_002")].mean == 94.1
    assert t1_map[("MAGE", "ARM_002")].sd == 18.5
    assert t1_map[("MAGE", "ARM_003")].mean == 89.1
    assert t1_map[("MAGE", "ARM_003")].sd == 29.7

    # eGFR: 82.6 (12.8), 76.5 (11.1), 80.7 (9.3)
    assert t1_map[("eGFR", "ARM_001")].mean == 82.6
    assert t1_map[("eGFR", "ARM_001")].sd == 12.8
    assert t1_map[("eGFR", "ARM_002")].mean == 76.5
    assert t1_map[("eGFR", "ARM_002")].sd == 11.1
    assert t1_map[("eGFR", "ARM_003")].mean == 80.7
    assert t1_map[("eGFR", "ARM_003")].sd == 9.3


def test_table2_insulin_gmean_gcv_and_ratios(rct4_extracted_data):
    """Verifies Errors 8, 15: Postprandial insulin reports gMean (% gCV) and multiplicative ratios."""
    outcomes = rct4_extracted_data.outcomes
    t2_ins = [o for o in outcomes if o.parameter_name == "Insulin AUC1-4h"]
    assert len(t2_ins) >= 9

    # Baseline geometric mean and % gCV
    t2_base = {o.arm_id: o for o in t2_ins if o.is_baseline}
    assert t2_base["ARM_001"].mean == 66.4
    assert t2_base["ARM_001"].dispersion_type == "gCV"
    assert t2_base["ARM_001"].metric_type == "geometric_mean_gcv"
    assert t2_base["ARM_002"].mean == 58.9
    assert t2_base["ARM_003"].mean == 65.3

    # Day 28 relative change ratio
    t2_d28 = {o.arm_id: o for o in t2_ins if str(o.timepoint_value) == "28"}
    assert t2_d28["ARM_001"].ratio == 1.0
    assert t2_d28["ARM_002"].ratio == 0.9
    assert t2_d28["ARM_003"].ratio == 0.8
    assert t2_d28["ARM_002"].effect_value == 0.9
    assert t2_d28["ARM_003"].effect_value == 0.8


def test_table3_urinary_8isopgf2a_se_dispersion(rct4_extracted_data):
    """Verifies Error 16: Urinary 8-iso-PGF2خ± reports Mean (SE) with Fasting and 24h separated."""
    outcomes = rct4_extracted_data.outcomes
    fasting_out = [o for o in outcomes if o.parameter_name == "8-iso-PGF2a (Fasting)"]
    h24_out = [o for o in outcomes if o.parameter_name == "8-iso-PGF2a (24h)"]
    assert len(fasting_out) == 6
    assert len(h24_out) == 6

    # Fasting baseline: 197.8 (27.0), 194.6 (29.4), 146.5 (18.5)
    f_base = {o.arm_id: o for o in fasting_out if o.is_baseline}
    assert f_base["ARM_001"].mean == 197.8
    assert f_base["ARM_001"].se == 27.0
    assert f_base["ARM_001"].sd is None
    assert f_base["ARM_001"].dispersion_type == "SE"
    assert f_base["ARM_002"].mean == 194.6
    assert f_base["ARM_002"].se == 29.4
    assert f_base["ARM_003"].mean == 146.5
    assert f_base["ARM_003"].se == 18.5


def test_table4_lipids_and_ketones_se_and_triglycerides(rct4_extracted_data):
    """Verifies Errors 9, 17: Table 4 lipids report Mean (SE), and Triglycerides reports log-scale gMean and ratio."""
    outcomes = rct4_extracted_data.outcomes
    tc_out = [o for o in outcomes if o.parameter_name == "Total Cholesterol"]
    assert len(tc_out) == 6
    tc_base = {o.arm_id: o for o in tc_out if o.is_baseline}
    assert tc_base["ARM_001"].mean == 212.8
    assert tc_base["ARM_001"].se == 7.7
    assert tc_base["ARM_001"].dispersion_type == "SE"
    assert tc_base["ARM_002"].mean == 200.1
    assert tc_base["ARM_002"].se == 8.4
    assert tc_base["ARM_003"].mean == 208.7
    assert tc_base["ARM_003"].se == 9.1

    # Triglycerides: baseline gMean (%CV) = 159.6, 131.5, 142.2; Day 29 ratio = 1.0, 0.8, 0.7
    tg_out = [o for o in outcomes if o.parameter_name == "Triglycerides"]
    tg_base = {o.arm_id: o for o in tg_out if o.is_baseline}
    assert tg_base["ARM_001"].mean == 159.6
    assert tg_base["ARM_001"].dispersion_type == "gCV"
    assert tg_base["ARM_002"].mean == 131.5
    assert tg_base["ARM_003"].mean == 142.2

    tg_d29 = {o.arm_id: o for o in tg_out if not o.is_baseline}
    assert tg_d29["ARM_001"].ratio == 1.0
    assert tg_d29["ARM_002"].ratio == 0.8
    assert tg_d29["ARM_003"].ratio == 0.7
    assert tg_d29["ARM_002"].effect_value == 0.8
    assert tg_d29["ARM_003"].effect_value == 0.7


def test_primary_endpoint_identification(rct4_extracted_data):
    """Verifies Error 14: PPG AUC1-4h at Day 28 is flagged is_primary=True and validated in QA."""
    outcomes = rct4_extracted_data.outcomes
    pri_out = [o for o in outcomes if o.is_primary]
    assert len(pri_out) == 3
    assert all(o.parameter_name == "PPG AUC1-4h" for o in pri_out)
    assert all(str(o.timepoint_value) == "28" for o in pri_out)

    pri_map = {o.arm_id: o for o in pri_out}
    assert pri_map["ARM_001"].change_value == -18.1
    assert pri_map["ARM_002"].change_value == -103.6
    assert pri_map["ARM_002"].effect_value == -85.5
    assert pri_map["ARM_003"].change_value == -122.9
    assert pri_map["ARM_003"].effect_value == -104.9

    # Check QA Inventory
    inv = rct4_extracted_data.inventory
    pri_inv = next(i for i in inv if "Primary Efficacy Endpoint" in i.item_label)
    assert pri_inv.extraction_status == "extracted_and_validated"
    assert pri_inv.recovered_record_count >= 3


def test_mage_and_hba1c_narrative_respective_tuples(rct4_extracted_data):
    """Verifies Errors 12, 13: MAGE and HbA1c respective tuples mapped to distinct arms without broadcast."""
    outcomes = rct4_extracted_data.outcomes
    mage_d1 = {o.arm_id: o for o in outcomes if o.parameter_name == "MAGE" and str(o.timepoint_value) == "1"}
    assert mage_d1["ARM_001"].change_value == 15.1
    assert mage_d1["ARM_001"].se == 3.5
    assert mage_d1["ARM_002"].change_value == 11.0
    assert mage_d1["ARM_002"].se == 3.7
    assert mage_d1["ARM_003"].change_value == 8.9
    assert mage_d1["ARM_003"].se == 3.7

    mage_d28 = {o.arm_id: o for o in outcomes if o.parameter_name == "MAGE" and str(o.timepoint_value) == "28"}
    assert mage_d28["ARM_001"].change_value == -4.7
    assert mage_d28["ARM_001"].se == 4.5
    assert mage_d28["ARM_002"].change_value == -3.7
    assert mage_d28["ARM_002"].se == 4.6
    assert mage_d28["ARM_003"].change_value == -2.2
    assert mage_d28["ARM_003"].se == 4.7

    hba1c_d29 = {o.arm_id: o for o in outcomes if o.parameter_name == "HbA1c" and str(o.timepoint_value) == "29"}
    assert hba1c_d29["ARM_001"].change_value == -0.11
    assert hba1c_d29["ARM_001"].se == 0.06
    assert hba1c_d29["ARM_002"].change_value == -0.46
    assert hba1c_d29["ARM_002"].se == 0.06
    assert hba1c_d29["ARM_002"].effect_value == -0.35
    assert hba1c_d29["ARM_003"].change_value == -0.63
    assert hba1c_d29["ARM_003"].se == 0.06
    assert hba1c_d29["ARM_003"].effect_value == -0.52


def test_safety_adverse_events(rct4_extracted_data):
    """Verifies Error 21: Full adverse events panel extracted."""
    outcomes = rct4_extracted_data.outcomes
    any_ae = {o.arm_id: o for o in outcomes if o.parameter_name == "Any Adverse Event"}
    assert any_ae["ARM_001"].events == 2
    assert any_ae["ARM_001"].percentage == 9.5
    assert any_ae["ARM_002"].events == 3
    assert any_ae["ARM_002"].percentage == 15.0
    assert any_ae["ARM_003"].events == 3
    assert any_ae["ARM_003"].percentage == 15.8

    gen_ae = {o.arm_id: o for o in outcomes if o.parameter_name == "Genital Infection"}
    assert gen_ae["ARM_001"].events == 0
    assert gen_ae["ARM_002"].events == 0
    assert gen_ae["ARM_003"].events == 1
    assert gen_ae["ARM_003"].percentage == 5.3


def test_excel_workbook_structure_and_columns(rct4_extracted_data):
    """Verifies Deliverables A, D: synchronized legacy + protocol sheets and updated column schemas."""
    file_path = "outputs/test_rct4_structure_export.xlsx"
    ExcelTrialExporter(rct4_extracted_data).export(file_path)
    assert os.path.exists(file_path), "Deliverable A workbook missing!"
    wb = openpyxl.load_workbook(file_path)
    expected_sheets = [
        "Master Excel Sheet", "Studies", "Arms", "Eligibility", "Outcomes",
        "Provenance", "Validation_Log", "Flattened_View", "Variables_Matrix", "QA_Audit_and_Inventory",
        "Treatment Arms", "Eligibility Criteria", "Baseline Characteristics", "Efficacy Outcomes",
        "Safety Outcomes", "Source Provenance", "Validation Log", "Quality Assessment"
    ]
    assert wb.sheetnames == expected_sheets
    assert wb.sheetnames[0] == "Master Excel Sheet"

    out_ws = wb["Outcomes"]
    headers = [cell.value for cell in out_ws[1]]
    for expected_col in ["se", "dispersion_type", "ratio", "is_primary", "analysis_population"]:
        assert expected_col in headers, f"Missing {expected_col} in Outcomes sheet!"


def test_error_report_completeness():
    """Verifies Deliverable B: Error report exists and has 21 validated error rows."""
    file_path = "outputs/RCT4_Extraction_Error_Report.xlsx"
    assert os.path.exists(file_path), "Deliverable B error report missing!"
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 22  # header + 21 errors
    error_ids = [r[0] for r in rows[1:]]
    for i in range(1, 22):
        assert f"ERR_{i:03d}" in error_ids
